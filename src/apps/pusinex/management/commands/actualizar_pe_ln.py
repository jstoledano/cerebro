import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone
from openpyxl import load_workbook

from apps.pusinex.models import (
    Distrito,
    DistritoLocal,
    Entidad,
    Localidad,
    Manzana,
    Municipio,
    PelnResidualSeccion,
    Seccion,
)


HEADERS = (
    "Entidad",
    "Distrito",
    "Municipio",
    "Sección",
    "Localidad",
    "Manzana",
    "Padrón",
    "Lista_nominal",
)
FILE_DATE_RE = re.compile(r"Corte_(\d{4}-\d{2}-\d{2})", re.IGNORECASE)
SHEET_DATE_RE = re.compile(
    r"Padron_lista_corte\s+(\d{4}-\d{2}-\d{2})",
    re.IGNORECASE,
)
BATCH_SIZE = 5000


class Command(BaseCommand):
    help = (
        "Actualiza Padrón Electoral y Lista Nominal desde todos los XLSX "
        "de un directorio y recalcula los agregados territoriales."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "ruta_directorio",
            type=str,
            help="Directorio que contiene los archivos XLSX de PE y LN",
        )

    def handle(self, *args, **options):
        directory = Path(options["ruta_directorio"]).expanduser().resolve()
        started_at = timezone.localtime()
        report = [
            "ACTUALIZACIÓN DE PADRÓN ELECTORAL Y LISTA NOMINAL",
            "=" * 60,
            f"Inicio: {started_at:%Y-%m-%d %H:%M:%S %Z}",
            f"Directorio: {directory}",
            "",
        ]
        incidences = []
        file_stats = []
        cut_date = None

        try:
            if not directory.is_dir():
                raise CommandError(f"El directorio no existe: {directory}")

            files = sorted(
                path
                for path in directory.iterdir()
                if path.is_file() and path.suffix.lower() == ".xlsx"
            )
            if not files:
                raise CommandError(
                    f"No se encontraron archivos XLSX en {directory}"
                )

            self.stdout.write(f"Archivos XLSX detectados: {len(files)}")
            cut_date = self._validate_cut_dates(files)
            rows, file_stats, incidences = self._read_files(files, cut_date)

            if not rows:
                raise CommandError("No hay registros válidos para procesar.")

            result = self._apply_rows(rows, cut_date, incidences)
            report.extend(
                self._success_report(
                    files,
                    cut_date,
                    file_stats,
                    result,
                    incidences,
                )
            )
            report_path = self._write_report(
                directory,
                cut_date,
                report,
            )

            self.stdout.write(self.style.SUCCESS("Actualización completada."))
            self.stdout.write(f"Fecha de corte: {cut_date:%Y-%m-%d}")
            self.stdout.write(
                f"Manzanas actualizadas: {result['blocks_updated']:,} | "
                f"Sin cambios: {result['blocks_unchanged']:,}"
            )
            self.stdout.write(
                "Localidades puntuales actualizadas: "
                f"{result['localities_updated']:,} | "
                f"Sin cambios: {result['localities_unchanged']:,}"
            )
            self.stdout.write(
                "Residuales de sección: "
                f"{result['residuals']:,} | "
                f"PE: {result['residual_pe']:,} | "
                f"LN: {result['residual_ln']:,}"
            )
            self.stdout.write(
                f"Incidencias: {len(incidences):,}"
            )
            self.stdout.write(
                self.style.SUCCESS(f"Reporte: {report_path}")
            )

        except Exception as exc:
            report.extend(
                [
                    "RESULTADO",
                    "-" * 60,
                    "ESTADO: ERROR",
                    (
                        "La transacción fue revertida o no llegó "
                        "a iniciarse."
                    ),
                    f"Detalle: {exc}",
                    "",
                ]
            )
            if file_stats:
                report.extend(self._format_file_stats(file_stats))
            if incidences:
                report.extend(self._format_incidences(incidences))

            report_path = None
            if directory.is_dir():
                report_path = self._write_report(
                    directory,
                    cut_date,
                    report,
                    error=True,
                )

            self.stderr.write(self.style.ERROR(str(exc)))
            if report_path:
                self.stderr.write(f"Reporte: {report_path}")

            if isinstance(exc, CommandError):
                raise
            raise CommandError(str(exc)) from exc

    def _validate_cut_dates(self, files):
        dates = set()
        details = []

        for path in files:
            match = FILE_DATE_RE.search(path.name)
            if not match:
                raise CommandError(
                    "El archivo no contiene la fecha de corte esperada "
                    f"en el nombre: {path.name}"
                )
            file_date = datetime.strptime(
                match.group(1),
                "%Y-%m-%d",
            ).date()

            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )
            try:
                sheet_dates = []
                for sheet_name in workbook.sheetnames:
                    sheet_match = SHEET_DATE_RE.search(sheet_name)
                    if sheet_match:
                        sheet_dates.append(
                            datetime.strptime(
                                sheet_match.group(1),
                                "%Y-%m-%d",
                            ).date()
                        )
            finally:
                workbook.close()

            if len(sheet_dates) != 1:
                raise CommandError(
                    f"{path.name}: se esperaba exactamente una hoja "
                    "con el patrón "
                    "'Padron_lista_corte YYYY-MM-DD'."
                )

            if sheet_dates[0] != file_date:
                raise CommandError(
                    f"{path.name}: la fecha del nombre "
                    f"({file_date}) difiere de la hoja "
                    f"({sheet_dates[0]})."
                )

            dates.add(file_date)
            details.append((path.name, file_date))

        if len(dates) != 1:
            rendered = ", ".join(
                f"{name}={date}"
                for name, date in details
            )
            raise CommandError(
                "Los archivos contienen fechas de corte distintas. "
                + rendered
            )

        return dates.pop()

    def _read_files(self, files, cut_date):
        rows = []
        incidences = []
        file_stats = []
        seen_edmslm = set()

        for path in files:
            workbook = load_workbook(
                path,
                read_only=True,
                data_only=True,
            )
            try:
                sheet_name = self._matching_sheet(
                    workbook,
                    cut_date,
                )
                worksheet = workbook[sheet_name]
                header_row = self._find_header_row(worksheet)
                stats = Counter(
                    read=0,
                    valid=0,
                    invalid=0,
                    duplicate=0,
                )

                iterator = worksheet.iter_rows(
                    min_row=header_row + 1,
                    min_col=1,
                    max_col=8,
                    values_only=True,
                )

                for row_number, values in enumerate(
                    iterator,
                    start=header_row + 1,
                ):
                    if all(value is None for value in values):
                        continue

                    stats["read"] += 1
                    try:
                        parsed = [
                            self._as_int(
                                value,
                                column,
                                row_number,
                            )
                            for value, column in zip(
                                values,
                                HEADERS,
                            )
                        ]
                    except ValueError as exc:
                        stats["invalid"] += 1
                        incidences.append(
                            {
                                "type": "FILA_INVALIDA",
                                "file": path.name,
                                "row": row_number,
                                "edmslm": self._partial_edmslm(
                                    values[:6]
                                ),
                                "detail": str(exc),
                            }
                        )
                        continue

                    (
                        entity,
                        district,
                        municipality,
                        section,
                        locality,
                        block,
                        pe,
                        ln,
                    ) = parsed

                    edmslm = self._edmslm(
                        entity,
                        district,
                        municipality,
                        section,
                        locality,
                        block,
                    )

                    if edmslm in seen_edmslm:
                        stats["duplicate"] += 1
                        incidences.append(
                            {
                                "type": "EDMSLM_DUPLICADO",
                                "file": path.name,
                                "row": row_number,
                                "edmslm": edmslm,
                                "detail": (
                                    "El EDMSLM ya apareció "
                                    "previamente en el lote."
                                ),
                            }
                        )
                        continue

                    seen_edmslm.add(edmslm)
                    rows.append(
                        {
                            "entity": entity,
                            "district": district,
                            "municipality": municipality,
                            "section": section,
                            "locality": locality,
                            "block": block,
                            "pe": pe,
                            "ln": ln,
                            "edmslm": edmslm,
                            "file": path.name,
                            "row": row_number,
                        }
                    )
                    stats["valid"] += 1

                file_stats.append(
                    {
                        "file": path.name,
                        "sheet": sheet_name,
                        **stats,
                    }
                )
            finally:
                workbook.close()

        return rows, file_stats, incidences

    def _apply_rows(self, rows, cut_date, incidences):
        sections = {
            section.seccion: section
            for section in Seccion.objects.select_related(
                "entidad",
                "distrito",
                "municipio",
            )
        }

        blocks = {
            (
                block.seccion_id,
                block.localidad,
                block.manzana,
            ): block
            for block in Manzana.objects.only(
                "id",
                "seccion_id",
                "localidad",
                "manzana",
                "padron",
                "lista_nominal",
            )
        }

        localities = {
            (
                locality.seccion_id,
                locality.localidad,
            ): locality
            for locality in Localidad.objects.only(
                "id",
                "seccion_id",
                "localidad",
                "padron",
                "lista_nominal",
            )
        }

        # Cada archivo mensual representa un snapshot completo del corte.
        # Conservamos los valores anteriores solo para clasificar cada fila
        # como actualizada o sin cambios antes de reiniciar PE/LN a cero.
        blocks_to_restore = []
        localities_to_restore = []
        residuals_to_create = []

        blocks_updated = 0
        blocks_unchanged = 0
        localities_updated = 0
        localities_unchanged = 0
        entity_ids = set()

        for row in rows:
            section = sections.get(row["section"])
            if section is None:
                incidences.append(
                    self._row_incidence(
                        row,
                        "SECCION_NO_ENCONTRADA",
                        (
                            "La sección indicada por el EDMSLM "
                            "no existe en la base de datos."
                        ),
                    )
                )
                continue

            if (
                section.entidad_id != row["entity"]
                or section.distrito_id != row["district"]
                or section.municipio_id != row["municipality"]
            ):
                incidences.append(
                    self._row_incidence(
                        row,
                        "EDMSLM_INCONSISTENTE",
                        (
                            "La entidad, distrito o municipio "
                            "del archivo no corresponde con "
                            "la sección registrada."
                        ),
                    )
                )
                continue

            # Manzana 9999 representa una localidad puntual.
            if row["block"] == 9999:
                locality = localities.get(
                    (
                        row["section"],
                        row["locality"],
                    )
                )

                if locality is None:
                    residuals_to_create.append(
                        PelnResidualSeccion(
                            seccion=section,
                            fecha_corte=cut_date,
                            localidad=row["locality"],
                            manzana=row["block"],
                            padron=row["pe"],
                            lista_nominal=row["ln"],
                            edmslm=row["edmslm"],
                            motivo="LOCALIDAD_NO_ENCONTRADA",
                            archivo=row["file"],
                            fila=row["row"],
                        )
                    )

                    incidences.append(
                        self._row_incidence(
                            row,
                            "PELN_RESIDUAL_SECCION",
                            (
                                "La localidad puntual no existe en la BGD. "
                                "PE/LN contabilizados a nivel de sección."
                            ),
                        )
                    )

                    entity_ids.add(row["entity"])
                    continue

                entity_ids.add(row["entity"])

                if (
                    locality.padron == row["pe"]
                    and locality.lista_nominal == row["ln"]
                ):
                    localities_unchanged += 1
                else:
                    localities_updated += 1

                locality.padron = row["pe"]
                locality.lista_nominal = row["ln"]
                localities_to_restore.append(locality)
                continue

            block = blocks.get(
                (
                    row["section"],
                    row["locality"],
                    row["block"],
                )
            )

            if block is None:
                residuals_to_create.append(
                    PelnResidualSeccion(
                        seccion=section,
                        fecha_corte=cut_date,
                        localidad=row["locality"],
                        manzana=row["block"],
                        padron=row["pe"],
                        lista_nominal=row["ln"],
                        edmslm=row["edmslm"],
                        motivo="MANZANA_NO_ENCONTRADA",
                        archivo=row["file"],
                        fila=row["row"],
                    )
                )

                incidences.append(
                    self._row_incidence(
                        row,
                        "PELN_RESIDUAL_SECCION",
                        (
                            "La manzana no existe en la BGD. "
                            "PE/LN contabilizados a nivel de sección."
                        ),
                    )
                )

                entity_ids.add(row["entity"])
                continue

            entity_ids.add(row["entity"])

            if (
                block.padron == row["pe"]
                and block.lista_nominal == row["ln"]
            ):
                blocks_unchanged += 1
            else:
                blocks_updated += 1

            block.padron = row["pe"]
            block.lista_nominal = row["ln"]
            blocks_to_restore.append(block)

        residual_pe = sum(item.padron for item in residuals_to_create)
        residual_ln = sum(item.lista_nominal for item in residuals_to_create)

        with transaction.atomic():
            # El corte mensual es un snapshot completo. Poner todo en cero
            # evita arrastrar PE/LN de manzanas o localidades ausentes del
            # archivo actual.
            Manzana.objects.update(
                padron=0,
                lista_nominal=0,
            )
            Localidad.objects.update(
                padron=0,
                lista_nominal=0,
            )

            # Restaurar todos los registros presentes en el corte, incluso
            # aquellos cuyo valor coincidía con el corte anterior.
            if blocks_to_restore:
                Manzana.objects.bulk_update(
                    blocks_to_restore,
                    ["padron", "lista_nominal"],
                    batch_size=BATCH_SIZE,
                )

            if localities_to_restore:
                Localidad.objects.bulk_update(
                    localities_to_restore,
                    ["padron", "lista_nominal"],
                    batch_size=BATCH_SIZE,
                )

            # El residual del mismo corte se reconstruye completamente en
            # cada ejecución para mantener el proceso determinista.
            PelnResidualSeccion.objects.filter(
                fecha_corte=cut_date
            ).delete()

            if residuals_to_create:
                PelnResidualSeccion.objects.bulk_create(
                    residuals_to_create,
                    batch_size=BATCH_SIZE,
                )

            totals = self._recalculate_aggregates(cut_date)

            if entity_ids:
                Entidad.objects.filter(
                    entidad__in=entity_ids
                ).update(fecha_corte_pe_ln=cut_date)

        return {
            "blocks_updated": blocks_updated,
            "blocks_unchanged": blocks_unchanged,
            "localities_updated": localities_updated,
            "localities_unchanged": localities_unchanged,
            "residuals": len(residuals_to_create),
            "residual_pe": residual_pe,
            "residual_ln": residual_ln,
            "updated": (
                blocks_updated
                + localities_updated
            ),
            "unchanged": (
                blocks_unchanged
                + localities_unchanged
            ),
            "entities": sorted(entity_ids),
            "totals": totals,
        }

    def _recalculate_aggregates(self, cut_date):
        block_totals = {
            item["seccion_id"]: item
            for item in Manzana.objects.values(
                "seccion_id"
            ).annotate(
                pe_sum=Sum("padron"),
                ln_sum=Sum("lista_nominal"),
            )
        }

        locality_totals = {
            item["seccion_id"]: item
            for item in Localidad.objects.values(
                "seccion_id"
            ).annotate(
                pe_sum=Sum("padron"),
                ln_sum=Sum("lista_nominal"),
            )
        }

        residual_totals = {
            item["seccion_id"]: item
            for item in PelnResidualSeccion.objects.filter(
                fecha_corte=cut_date
            ).values(
                "seccion_id"
            ).annotate(
                pe_sum=Sum("padron"),
                ln_sum=Sum("lista_nominal"),
            )
        }

        sections = list(Seccion.objects.all())

        for section in sections:
            block_data = block_totals.get(section.seccion)
            locality_data = locality_totals.get(section.seccion)
            residual_data = residual_totals.get(section.seccion)

            section.pe = (
                (block_data["pe_sum"] if block_data else 0)
                + (locality_data["pe_sum"] if locality_data else 0)
                + (residual_data["pe_sum"] if residual_data else 0)
            )
            section.ln = (
                (block_data["ln_sum"] if block_data else 0)
                + (locality_data["ln_sum"] if locality_data else 0)
                + (residual_data["ln_sum"] if residual_data else 0)
            )

        if sections:
            Seccion.objects.bulk_update(
                sections,
                ["pe", "ln"],
                batch_size=BATCH_SIZE,
            )

        municipality_totals = {
            item["municipio_id"]: item
            for item in Seccion.objects.values(
                "municipio_id"
            ).annotate(
                pe_sum=Sum("pe"),
                ln_sum=Sum("ln"),
            )
        }
        municipalities = list(Municipio.objects.all())
        for municipality in municipalities:
            data = municipality_totals.get(
                municipality.municipio
            )
            municipality.pe = data["pe_sum"] if data else 0
            municipality.ln = data["ln_sum"] if data else 0
        if municipalities:
            Municipio.objects.bulk_update(
                municipalities,
                ["pe", "ln"],
                batch_size=BATCH_SIZE,
            )

        district_totals = {
            item["distrito_id"]: item
            for item in Seccion.objects.values(
                "distrito_id"
            ).annotate(
                pe_sum=Sum("pe"),
                ln_sum=Sum("ln"),
            )
        }
        districts = list(Distrito.objects.all())
        for district in districts:
            data = district_totals.get(
                district.distrito
            )
            district.pe = data["pe_sum"] if data else 0
            district.ln = data["ln_sum"] if data else 0
        if districts:
            Distrito.objects.bulk_update(
                districts,
                ["pe", "ln"],
                batch_size=BATCH_SIZE,
            )

        local_totals = {
            item["distrito_l_id"]: item
            for item in Seccion.objects.values(
                "distrito_l_id"
            ).annotate(
                pe_sum=Sum("pe"),
                ln_sum=Sum("ln"),
            )
        }
        local_districts = list(DistritoLocal.objects.all())
        for district in local_districts:
            data = local_totals.get(
                district.distrito_local
            )
            district.pe = data["pe_sum"] if data else 0
            district.ln = data["ln_sum"] if data else 0
        if local_districts:
            DistritoLocal.objects.bulk_update(
                local_districts,
                ["pe", "ln"],
                batch_size=BATCH_SIZE,
            )

        entity_totals = {
            item["entidad_id"]: item
            for item in Seccion.objects.values(
                "entidad_id"
            ).annotate(
                pe_sum=Sum("pe"),
                ln_sum=Sum("ln"),
            )
        }
        entities = list(Entidad.objects.all())
        for entity in entities:
            data = entity_totals.get(entity.entidad)
            entity.pe = data["pe_sum"] if data else 0
            entity.ln = data["ln_sum"] if data else 0
        if entities:
            Entidad.objects.bulk_update(
                entities,
                ["pe", "ln"],
                batch_size=BATCH_SIZE,
            )

        return {
            "sections": len(sections),
            "municipalities": len(municipalities),
            "districts": len(districts),
            "local_districts": len(local_districts),
            "entities": len(entities),
        }

    def _success_report(
        self,
        files,
        cut_date,
        file_stats,
        result,
        incidences,
    ):
        totals = result["totals"]
        lines = [
            "RESULTADO",
            "-" * 60,
            "ESTADO: COMPLETADO",
            f"Fecha de corte: {cut_date:%Y-%m-%d}",
            f"Archivos procesados: {len(files)}",
            f"Manzanas actualizadas: {result['blocks_updated']}",
            f"Manzanas sin cambios: {result['blocks_unchanged']}",
            (
                "Localidades puntuales actualizadas: "
                f"{result['localities_updated']}"
            ),
            (
                "Localidades puntuales sin cambios: "
                f"{result['localities_unchanged']}"
            ),
            f"Residuales de sección: {result['residuals']}",
            f"PE residual: {result['residual_pe']}",
            f"LN residual: {result['residual_ln']}",
            f"Registros actualizados: {result['updated']}",
            f"Registros sin cambios: {result['unchanged']}",
            f"Incidencias: {len(incidences)}",
            "",
            "AGREGADOS RECALCULADOS",
            "-" * 60,
            f"Secciones: {totals['sections']}",
            f"Municipios: {totals['municipalities']}",
            f"Distritos federales: {totals['districts']}",
            f"Distritos locales: {totals['local_districts']}",
            f"Entidades: {totals['entities']}",
            "",
        ]
        lines.extend(self._format_file_stats(file_stats))
        lines.extend(self._format_incidences(incidences))
        return lines

    def _format_file_stats(self, file_stats):
        lines = ["RESUMEN POR ARCHIVO", "-" * 60]
        for stats in file_stats:
            lines.extend(
                [
                    f"Archivo: {stats['file']}",
                    f"  Hoja: {stats['sheet']}",
                    f"  Filas leídas: {stats['read']}",
                    f"  Filas válidas: {stats['valid']}",
                    f"  Filas inválidas: {stats['invalid']}",
                    (
                        "  EDMSLM duplicados: "
                        f"{stats['duplicate']}"
                    ),
                ]
            )
        lines.append("")
        return lines

    def _format_incidences(self, incidences):
        lines = ["INCIDENCIAS", "-" * 60]
        if not incidences:
            return lines + ["Sin incidencias.", ""]

        counts = Counter(
            item["type"]
            for item in incidences
        )
        lines.append("Resumen:")
        for kind in sorted(counts):
            lines.append(f"  {kind}: {counts[kind]}")
        lines.extend(["", "Detalle:"])

        for item in incidences:
            lines.append(
                f"[{item['type']}] "
                f"{item['edmslm']} | "
                f"{item['file']}:{item['row']} | "
                f"{item['detail']}"
            )
        lines.append("")
        return lines

    def _write_report(
        self,
        directory,
        cut_date,
        report,
        error=False,
    ):
        now = timezone.localtime()
        date_part = (
            cut_date.isoformat()
            if cut_date
            else "sin_fecha"
        )
        error_part = "error_" if error else ""
        filename = (
            f"reporte_pe_ln_{error_part}{date_part}_"
            f"{now:%Y%m%d-%H%M%S}.txt"
        )
        path = directory / filename
        report.extend(
            [
                f"Fin: {now:%Y-%m-%d %H:%M:%S %Z}",
                "",
            ]
        )
        path.write_text(
            "\n".join(report),
            encoding="utf-8",
        )
        return path

    def _matching_sheet(self, workbook, cut_date):
        for sheet_name in workbook.sheetnames:
            match = SHEET_DATE_RE.search(sheet_name)
            if (
                match
                and datetime.strptime(
                    match.group(1),
                    "%Y-%m-%d",
                ).date()
                == cut_date
            ):
                return sheet_name
        raise CommandError(
            "No se encontró la hoja correspondiente "
            f"al corte {cut_date}."
        )

    def _find_header_row(self, worksheet):
        limit = min(10, worksheet.max_row)
        for row_number in range(1, limit + 1):
            values = tuple(
                worksheet.cell(
                    row=row_number,
                    column=column,
                ).value
                for column in range(1, 9)
            )
            if values == HEADERS:
                return row_number

        raise CommandError(
            f"{worksheet.title}: no se encontró el "
            "encabezado esperado en las primeras 10 filas."
        )

    def _as_int(self, value, column, row_number):
        if value is None or value == "":
            raise ValueError(
                f"Fila {row_number}: {column} está vacío."
            )
        if isinstance(value, bool):
            raise ValueError(
                f"Fila {row_number}: {column} "
                "contiene un booleano."
            )
        if (
            isinstance(value, float)
            and not value.is_integer()
        ):
            raise ValueError(
                f"Fila {row_number}: {column} "
                f"contiene un decimal: {value}."
            )
        try:
            parsed = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"Fila {row_number}: {column} "
                f"no es un entero válido: {value!r}."
            ) from exc
        if parsed < 0:
            raise ValueError(
                f"Fila {row_number}: {column} "
                "contiene un valor negativo."
            )
        return parsed

    def _row_incidence(self, row, kind, detail):
        return {
            "type": kind,
            "file": row["file"],
            "row": row["row"],
            "edmslm": row["edmslm"],
            "detail": detail,
        }

    def _edmslm(
        self,
        entity,
        district,
        municipality,
        section,
        locality,
        block,
    ):
        return (
            f"{entity:02d}"
            f"{district:02d}"
            f"{municipality:03d}"
            f"{section:04d}"
            f"{locality:04d}"
            f"{block:04d}"
        )

    def _partial_edmslm(self, values):
        widths = (2, 2, 3, 4, 4, 4)
        parts = []
        for value, width in zip(values, widths):
            try:
                parts.append(
                    f"{int(value):0{width}d}"
                )
            except (TypeError, ValueError):
                parts.append("?" * width)
        return "".join(parts)
